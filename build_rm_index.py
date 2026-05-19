"""
Rebuild data/rm_index.json from the authoritative raw-material exports.

Inputs (snapshot in data/sources/):
  - Items863.xls               — RM master list (SpreadsheetML XML), 237 RMs
  - RMsResults_with_CAS.xlsx   — RM → CAS + Concentration %

Output:
  - data/rm_index.json         — { RM#: {cas, cas_all, name, synonyms,
                                   active_fraction, strength_pct,
                                   strength_source, needs_review} }
  - data/sources/rm_gap_report.txt — missing-CAS / dirty-data triage list

This step is DATA ONLY: supplied strength is parsed (mainly from the trade
name, cross-checked against the Concentration column) and stored as
`active_fraction` for a later classification-conversion step. It does NOT
change GHS classification behaviour.

Run:  python3 build_rm_index.py
"""

from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

BASE = Path(__file__).parent
SRC = BASE / "data" / "sources"
MASTER_XLS = SRC / "Items863.xls"
CAS_XLSX = SRC / "RMsResults_with_CAS.xlsx"
OUT = BASE / "data" / "rm_index.json"
REPORT = SRC / "rm_gap_report.txt"

_NS = {"s": "urn:schemas-microsoft-com:office:spreadsheet"}
_RM_RE = re.compile(r"RM\d+(?:\.\d+)?", re.I)
_PCT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*%")
_CAS_RE = re.compile(r"\b\d{2,7}-\d{2}-\d\b")


def _norm_rm(raw: str) -> str:
    """'RM0002 - HEDP 60%' / 'RM0001EPA' → 'RM0002' / 'RM0001'."""
    m = _RM_RE.search(str(raw or "").upper())
    return m.group(0) if m else str(raw or "").strip().upper()


def _strength_from_name(name: str) -> float | None:
    """First percentage in the trade name, e.g. 'Sulfuric Acid 50%' → 50.0."""
    m = _PCT_RE.search(name or "")
    if not m:
        # 'Glacial' acetic acid and similar neat liquids are effectively 100 %.
        if re.search(r"\bglacial\b", name or "", re.I):
            return 100.0
        return None
    try:
        v = float(m.group(1))
        return v if 0 < v <= 100 else None
    except ValueError:
        return None


def _read_master() -> dict[str, dict]:
    root = ET.parse(MASTER_XLS).getroot()

    def cells(r):
        out = []
        for c in r.findall("s:Cell", _NS):
            d = c.find("s:Data", _NS)
            out.append((d.text or "").strip() if d is not None else "")
        return out

    rows = [cells(x) for x in root.findall(".//s:Table/s:Row", _NS)][1:]
    master: dict[str, dict] = {}
    for d in rows:
        if not d or not d[0]:
            continue
        rm = _norm_rm(d[0])
        master.setdefault(rm, {
            "name": (d[1] if len(d) > 1 else "") or rm,
            "synonyms": d[2] if len(d) > 2 else "",
        })
    return master


def _read_cas() -> dict[str, dict]:
    wb = openpyxl.load_workbook(CAS_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out: dict[str, dict] = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r or not r[0]:
            continue
        rm = _norm_rm(r[0])
        cas_raw = str(r[5] or "").strip() if len(r) > 5 else ""
        cas_all = _CAS_RE.findall(cas_raw)
        conc = r[6] if len(r) > 6 else None
        out[rm] = {
            "name": (r[1] or "").strip() if len(r) > 1 else "",
            "synonyms": (r[2] or "").strip() if len(r) > 2 else "",
            "cas": cas_all[0] if cas_all else "",
            "cas_all": cas_all,
            "conc": conc if isinstance(conc, (int, float)) else None,
        }
    return out


def main() -> None:
    master = _read_master()
    casmap = _read_cas()
    all_rms = sorted(set(master) | set(casmap))

    index: dict[str, dict] = {}
    missing_cas, multi_cas, conc_conflict, dirty_conc, cas_only = [], [], [], [], []

    for rm in all_rms:
        m = master.get(rm, {})
        c = casmap.get(rm, {})
        name = m.get("name") or c.get("name") or rm
        syn = m.get("synonyms") or c.get("synonyms") or ""
        cas = c.get("cas", "")
        cas_all = c.get("cas_all", [])

        name_pct = _strength_from_name(name)
        col_pct = c.get("conc")
        col_ok = isinstance(col_pct, (int, float)) and 0 < col_pct <= 100

        if name_pct is not None:
            strength, src = name_pct, "name"
            if col_ok and abs(col_pct - name_pct) > 2:
                conc_conflict.append((rm, name, name_pct, col_pct))
        elif col_ok:
            strength, src = float(col_pct), "column"
        else:
            strength, src = 100.0, "default"
        if col_pct is not None and not col_ok:
            dirty_conc.append((rm, name, col_pct))

        needs_review = (not cas) or src == "default" or bool(conc_conflict
                        and conc_conflict[-1][0] == rm)
        if not cas:
            missing_cas.append((rm, name))
        if len(cas_all) > 1:
            multi_cas.append((rm, name, cas_all))
        if rm not in master and rm in casmap:
            cas_only.append(rm)

        index[rm] = {
            "cas": cas,
            "cas_all": cas_all,
            "name": name,
            "synonyms": syn,
            "active_fraction": round(strength / 100.0, 4),
            "strength_pct": strength,
            "strength_source": src,
            "needs_review": needs_review,
        }

    OUT.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = []
    lines.append(f"RM index rebuilt: {len(index)} RMs "
                 f"(master {len(master)}, cas-table {len(casmap)})")
    lines.append(f"WITH cas: {sum(1 for v in index.values() if v['cas'])} | "
                 f"MISSING cas: {len(missing_cas)}")
    lines.append(f"strength source — name: "
                 f"{sum(1 for v in index.values() if v['strength_source']=='name')}, "
                 f"column: {sum(1 for v in index.values() if v['strength_source']=='column')}, "
                 f"default(100%): {sum(1 for v in index.values() if v['strength_source']=='default')}")
    lines.append(f"multi-CAS: {len(multi_cas)} | conc name/col conflicts: "
                 f"{len(conc_conflict)} | dirty conc values: {len(dirty_conc)} | "
                 f"cas-table-only RMs: {len(cas_only)}")
    lines.append("")
    lines.append("== MISSING CAS (need lookup) ==")
    lines += [f"  {rm}  {nm}" for rm, nm in sorted(missing_cas)]
    lines.append("")
    lines.append("== MULTI-CAS (mixtures/polymers — primary = first) ==")
    lines += [f"  {rm}  {nm}  {ca}" for rm, nm, ca in multi_cas]
    lines.append("")
    lines.append("== NAME vs COLUMN concentration conflict (>2%) ==")
    lines += [f"  {rm}  {nm}  name={a}  col={b}" for rm, nm, a, b in conc_conflict]
    lines.append("")
    lines.append("== DIRTY concentration column value (ignored) ==")
    lines += [f"  {rm}  {nm}  conc={cv}" for rm, nm, cv in dirty_conc]
    REPORT.write_text("\n".join(lines), encoding="utf-8")

    print("\n".join(lines[:4]))
    print(f"\nWrote {OUT}  and  {REPORT}")


if __name__ == "__main__":
    main()
